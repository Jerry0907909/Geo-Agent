"""Word 文档解析工具"""

import logging
from typing import Union
from io import BytesIO

logger = logging.getLogger(__name__)


def parse_doc_content(content: Union[bytes, str], file_ext: str = ".docx") -> str:
    """解析 Word 文档内容
    
    Args:
        content: 文件内容（bytes）或文件路径（str）
        file_ext: 文件扩展名 (.docx 或 .doc)
        
    Returns:
        提取的文本内容
    """
    if file_ext.lower() == ".docx":
        return parse_docx_content(content)
    else:
        return parse_doc_legacy_content(content)


def parse_docx_content(content: Union[bytes, str]) -> str:
    """解析 .docx 文件"""
    try:
        from docx import Document
        
        if isinstance(content, str):
            doc = Document(content)
        else:
            doc = Document(BytesIO(content))
        
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        
        # 也提取表格内容
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text_parts.append(" | ".join(row_text))
        
        return "\n".join(text_parts)
        
    except ImportError:
        logger.error("python-docx 未安装")
        raise ImportError("请安装 python-docx: pip install python-docx")
    except Exception as e:
        logger.error(f"DOCX 解析失败: {e}")
        raise


def parse_doc_legacy_content(content: Union[bytes, str]) -> str:
    """解析旧版 .doc 文件"""
    try:
        import textract
        
        if isinstance(content, bytes):
            # textract 需要文件路径，创建临时文件
            import tempfile
            import os
            
            with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp:
                tmp.write(content)
                tmp_path = tmp.name
            
            try:
                text = textract.process(tmp_path).decode('utf-8', errors='ignore')
            finally:
                os.unlink(tmp_path)
            
            return text
        else:
            return textract.process(content).decode('utf-8', errors='ignore')
            
    except ImportError:
        logger.warning("textract 未安装，尝试使用 antiword")
        # 如果是 bytes，无法处理
        if isinstance(content, bytes):
            raise ImportError("请安装 textract 来处理 .doc 文件: pip install textract")
        
        import subprocess
        try:
            result = subprocess.run(
                ['antiword', content],
                capture_output=True,
                text=True
            )
            return result.stdout
        except FileNotFoundError:
            raise ImportError("请安装 antiword 或 textract 来处理 .doc 文件")
    except Exception as e:
        logger.error(f"DOC 解析失败: {e}")
        raise
